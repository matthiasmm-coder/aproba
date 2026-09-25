#!/usr/bin/env python3
"""PERFIL de un archivo de migración (Excel o CSV) — LECTURA SOLA: nada sale del ordenador.

Fase 3 de web/MIGRACION.md: antes de preparar nada, saber QUÉ trae el archivo del cliente.
Por hoja: fila de cabecera, nº de filas, columnas (tipo detectado y relleno); personas
(NIE/DNI, pasaportes) y empresas (CIF válidos); la MISMA persona en varias filas (típico de un
listado de facturas); homónimos con documentos distintos o casi iguales (errores de tecleo);
importes y estado del cobro → TOTALES DE CONTROL (para verificar.ts); columnas SENSIBLES
(anotaciones internas de cobro: nunca se importan); marcas de pago («Primer pago (1-2)»).

Uso:
  python3 scripts/migracion/perfilar.py <archivo.xlsx|.csv> [--hoja NOMBRE] [--cabecera N]
          [--salida perfil.md] [--control control.json]
"""
import argparse, csv, datetime as dt, json, re, sys, unicodedata
from collections import Counter, defaultdict

RE_NIE = re.compile(r"^[XYZ]\d{7}[A-Z]$")
RE_DNI = re.compile(r"^\d{8}[A-Z]$")
RE_CIF = re.compile(r"^[ABCDEFGHJNPQRSUVW]\d{7}[0-9A-J]$")
RE_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", re.I)
RE_TEL = re.compile(r"^\+?[\d\s().-]{9,}$")
RE_FECHA = re.compile(r"^\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}$|^\d{4}-\d{2}-\d{2}")
RE_PASAPORTE = re.compile(r"^[A-Z0-9]{6,12}$")
RE_PAGO = re.compile(r"\(\s*\d\s*[-/]\s*\d\s*\)|factura\s+\d\s+de\s+\d|(primer|segundo|tercer|1er|[123][ºªo])\s+pago", re.I)
# Cabeceras que suelen llevar control INTERNO de cobros (dinero fuera de factura): nunca se importan.
RE_SENSIBLE = re.compile(r"observ|falta|coment|^\s*b\s*$|\ben b\b|\(b\)|negro|efectivo|a cuenta|sin factura|caja", re.I)
# Columnas de documento de identidad por su cabecera (antes que por sus valores: un nº de
# factura «AGC0122.2026» parece un pasaporte).
RE_CAB_DOC = re.compile(r"\bnie\b|\bnif\b|\bdni\b|\bcif\b|documento|pasaporte|identificaci", re.I)
RE_CAB_NO_DOC = re.compile(r"factura|referencia|\bref\b|expediente|n[ºo°]\s*$", re.I)
RE_ESTADO_COBRO = re.compile(r"cobr|pagad|pendiente|impagad", re.I)
RE_PENDIENTE = re.compile(r"\b(no|pendiente|pte|parcial|impagad|debe|sin cobrar|por cobrar)", re.I)


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def doc(v):
    return re.sub(r"[\s.-]", "", str(v or "")).upper()


def cif_valido(v):
    """CIF español con dígito de control (evita confundir un pasaporte con una empresa)."""
    v = doc(v)
    if not RE_CIF.match(v):
        return False
    digitos = v[1:8]
    pares = sum(int(digitos[i]) for i in (1, 3, 5))
    impares = sum(sum(divmod(int(digitos[i]) * 2, 10)) for i in (0, 2, 4, 6))
    control = (10 - (pares + impares) % 10) % 10
    return v[8] in (str(control), "JABCDEFGHI"[control])


def importe(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.,-]", "", str(v))
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def tipo_de(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return "fecha"
    if isinstance(v, (int, float)):
        return "número"
    s = str(v).strip()
    d = doc(s)
    if RE_NIE.match(d) or RE_DNI.match(d):
        return "NIE/DNI"
    if cif_valido(d):
        return "CIF"
    if RE_EMAIL.match(s):
        return "email"
    if RE_FECHA.match(s):
        return "fecha"
    if importe(s) is not None and re.fullmatch(r"[\d.,\s€-]+", s):
        return "número"
    if RE_TEL.match(s):
        return "teléfono"
    if RE_PASAPORTE.match(d) and re.search(r"\d", d) and re.search(r"[A-Z]", d):
        return "pasaporte?"
    return "texto"


def leer(ruta, hoja):
    if ruta.lower().endswith(".csv"):
        with open(ruta, newline="", encoding="utf-8-sig") as f:
            muestra = f.read(4096)
            f.seek(0)
            dialecto = csv.Sniffer().sniff(muestra, delimiters=";,\t")
            return {"CSV": [list(r) for r in csv.reader(f, dialecto)]}
    import openpyxl, warnings
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    hojas = [hoja] if hoja else wb.sheetnames
    return {h: [list(r) for r in wb[h].iter_rows(values_only=True)] for h in hojas}


def eur(n):
    return f"{n:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


def significativa(r):
    """Una fila de datos de verdad: 3 celdas o más con algo que no sea un cero de fórmula."""
    return sum(1 for c in r if c not in (None, "") and not (isinstance(c, (int, float)) and c == 0)) >= 3


def fila_cabecera(filas):
    """La fila con más celdas de TEXTO entre las 30 primeras, seguida de filas con datos."""
    mejor, puntos = 0, -1
    for i, r in enumerate(filas[:30]):
        textos = sum(1 for c in r if isinstance(c, str) and c.strip() and not re.fullmatch(r"[\d.,\s€%-]+", c))
        datos_despues = sum(1 for rr in filas[i + 1:i + 6] if sum(1 for c in rr if c not in (None, "")) >= max(2, textos // 2))
        p = textos + datos_despues
        if textos >= 3 and p > puntos:
            mejor, puntos = i, p
    return mejor


def perfil_hoja(nombre, filas, cab_forzada, salida):
    """Escribe el perfil de la hoja y devuelve sus totales de control (o None)."""
    control = {}
    h = cab_forzada - 1 if cab_forzada else fila_cabecera(filas)
    cab = [str(c).strip() if c is not None else "" for c in (filas[h] if filas else [])]
    datos = [r for r in filas[h + 1:] if significativa(r)]
    ncol = max([len(cab)] + [len(r) for r in datos[:200]]) if datos else len(cab)
    cab += [""] * (ncol - len(cab))
    # Cabecera en DOS filas (típica de un Excel hecho a mano): los huecos, de la fila de arriba.
    if h > 0:
        arriba = [str(c).strip() if isinstance(c, str) else "" for c in filas[h - 1]] + [""] * ncol
        if sum(1 for c in arriba if c) >= 3:
            cab = [c or arriba[j] for j, c in enumerate(cab)]
    p = salida.append
    p(f"\n## Hoja «{nombre}» — cabecera en la fila {h + 1}, {len(datos)} filas de datos\n")
    if not datos:
        p("_Sin datos._")
        return None

    # Columnas: tipo mayoritario, relleno, ejemplos NO personales solo para texto corto.
    tipos = {}
    p("| # | Columna | Tipo | Relleno | Distintos |")
    p("|---|---|---|---|---|")
    for j in range(ncol):
        vals = [r[j] for r in datos if j < len(r) and r[j] not in (None, "")]
        t = Counter(tipo_de(v) for v in vals).most_common(1)[0][0] if vals else "vacía"
        tipos[j] = t
        p(f"| {j + 1} | {cab[j] or '(sin nombre)'} | {t} | {round(100 * len(vals) / len(datos))} % | {len(set(map(str, vals)))} |")

    # Columnas sensibles (solo el nombre y cuántas filas: nunca el contenido).
    sens = [j for j in range(ncol) if RE_SENSIBLE.search(cab[j] or "")]
    if sens:
        p("\n**⚠ Columnas sensibles — NO se importan** (control interno de cobros: una migración nunca guarda dinero fuera de factura):")
        for j in sens:
            n = sum(1 for r in datos if j < len(r) and r[j] not in (None, ""))
            p(f"- «{cab[j]}»: {n} filas con contenido")

    # Personas y empresas.
    col_doc = [j for j in range(ncol) if RE_CAB_DOC.search(cab[j] or "") and not RE_CAB_NO_DOC.search(cab[j] or "")] \
        or [j for j, t in tipos.items() if t in ("NIE/DNI", "CIF") and not RE_CAB_NO_DOC.search(cab[j] or "")]
    col_nombre = next((j for j in range(ncol) if re.search(r"nombre|cliente|titular|raz[oó]n", cab[j] or "", re.I) and tipos[j] == "texto"), None)
    por_doc = defaultdict(list)
    empresas = set()
    for i, r in enumerate(datos):
        for j in col_doc:
            v = doc(r[j]) if j < len(r) and r[j] not in (None, "") else ""
            if not v:
                continue
            if cif_valido(v):
                empresas.add(v)
            else:
                por_doc[v].append(i)
            break
    varias = {d: ix for d, ix in por_doc.items() if len(ix) > 1}
    p(f"\n**Personas**: {len(por_doc)} documentos distintos · **empresas** (CIF válido): {len(empresas)}")
    if varias:
        dist = Counter(len(ix) for ix in varias.values())
        p(f"- {len(varias)} personas aparecen en varias filas ({', '.join(f'{n} filas: {c}' for n, c in sorted(dist.items()))}) → probablemente un **listado de facturas**: cada fila es un servicio, la persona se crea una vez.")
    sin_doc = sum(1 for r in datos if not any(j < len(r) and r[j] not in (None, "") for j in col_doc))
    p(f"- {sin_doc} filas sin documento (se identificarán por nombre: revisar homónimos)")
    if col_nombre is not None:
        docs_de = defaultdict(set)
        for r in datos:
            n = norm(r[col_nombre]) if col_nombre < len(r) else ""
            d = next((doc(r[j]) for j in col_doc if j < len(r) and r[j] not in (None, "") and not cif_valido(r[j])), "")
            if n and d:
                docs_de[n].add(d)
        homonimos = {n: ds for n, ds in docs_de.items() if len(ds) > 1}
        casi = [(n, sorted(ds)) for n, ds in homonimos.items() if any(sum(a != b for a, b in zip(x, y)) <= 1 and len(x) == len(y) for x in ds for y in ds if x < y)]
        if homonimos:
            p(f"- {len(homonimos)} nombres con varios documentos → dos personas, o un error de tecleo: **decidir a mano** ({len(casi)} difieren en un solo carácter{': ' + '; '.join(f'{n.upper()} {ds}' for n, ds in casi[:6]) if casi else ''})")

    # Importes y cobro → totales de control.
    col_imp = [j for j, t in tipos.items() if t == "número" and re.search(r"total|importe|b\.?i|base|iva|precio|€", cab[j] or "", re.I)]
    col_cobro = next((j for j in range(ncol) if RE_ESTADO_COBRO.search(cab[j] or "") and tipos[j] == "texto"), None)
    if col_imp:
        p("\n**Importes** (totales de control):")
        for j in col_imp:
            vals = [importe(r[j]) for r in datos if j < len(r)]
            vals = [v for v in vals if v is not None]
            p(f"- «{cab[j]}»: {len(vals)} valores · total {eur(sum(vals))}")
    if col_cobro is not None:
        estados = Counter(str(r[col_cobro]).strip().upper() for r in datos if col_cobro < len(r) and r[col_cobro] not in (None, ""))
        p(f"- «{cab[col_cobro]}»: " + " · ".join(f"{k} {v}" for k, v in estados.most_common(8)))
    j_total = next((j for j in col_imp if re.search(r"total", cab[j] or "", re.I)), col_imp[-1] if col_imp else None)
    if j_total is not None:
        tot = [importe(r[j_total]) for r in datos if j_total < len(r)]
        pend = [importe(r[j_total]) for r in datos if col_cobro is not None and col_cobro < len(r) and RE_PENDIENTE.search(str(r[col_cobro] or "")) and j_total < len(r)]
        control.update({
            "hoja": nombre, "columnaImporte": cab[j_total], "facturas": sum(1 for v in tot if v is not None),
            "importeTotal": round(sum(v for v in tot if v is not None), 2),
            "pendientes": sum(1 for v in pend if v is not None), "importePendiente": round(sum(v for v in pend if v is not None), 2),
            "personas": len(por_doc), "empresas": len(empresas),
        })

    # Marcas de pago en columnas de texto.
    marcas = sum(1 for r in datos if any(isinstance(c, str) and RE_PAGO.search(c) for c in r))
    if marcas:
        p(f"\n**Pagos fraccionados**: {marcas} filas con «primer/segundo pago», «(1-2)» o «factura 1 de 2» → el motor junta las facturas de un mismo servicio (misma persona, mismo servicio, mismo concepto).")
    return control or None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("archivo")
    ap.add_argument("--hoja")
    ap.add_argument("--cabecera", type=int, help="nº de fila (1 = primera) de la cabecera, si la detección falla")
    ap.add_argument("--salida", help="guardar el informe (Markdown) en este archivo")
    ap.add_argument("--control", help="guardar los totales de control (JSON) para verificar.ts")
    a = ap.parse_args()
    salida = [f"# Perfil de «{a.archivo.split('/')[-1]}» — {dt.date.today():%d/%m/%Y}"]
    candidatos = [c for nombre, filas in leer(a.archivo, a.hoja).items() if (c := perfil_hoja(nombre, filas, a.cabecera, salida))]
    # Totales de control: los de la hoja con más facturas (o la elegida con --hoja).
    control = max(candidatos, key=lambda c: c["facturas"]) if candidatos else {}
    texto = "\n".join(salida) + "\n"
    if a.salida:
        open(a.salida, "w", encoding="utf-8").write(texto)
        print(f"informe: {a.salida}")
    else:
        sys.stdout.write(texto)
    if a.control:
        json.dump(control, open(a.control, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print(f"totales de control: {a.control} (revísalos: la hoja y la columna del importe)")


if __name__ == "__main__":
    main()
